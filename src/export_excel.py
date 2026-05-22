from __future__ import annotations

import argparse
import base64
import csv
import json
import math
import mimetypes
import re
import shutil
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urljoin, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from PIL import Image as PILImage

try:
    from lxml import html as lxml_html
except ImportError:  # pragma: no cover
    lxml_html = None

try:
    from tqdm import tqdm
except ImportError:  # pragma: no cover
    tqdm = None


SOURCE_URL = "https://oge.fipi.ru/bank/index.php?proj=DE0E276E497AB3784C3FC4CC20248DC0"
TITLE = "ФИПИ ОГЭ Математика — полный экспорт заданий"
SUMMARY_SHEET = "Общая информация"
INDEX_SHEET = "Перечень заданий"
TASKS_SHEET = "Задания"
ERRORS_SHEET = "Ошибки"
MAX_IMAGE_WIDTH_PX = 680
MAX_CELL_TEXT = 32000


@dataclass
class ExportStats:
    tasks_processed: int = 0
    cards_created: int = 0
    images_embedded: int = 0
    images_linked: int = 0
    rendered_png_created: int = 0
    rendered_fallbacks: int = 0
    render_resources_found: int = 0
    render_resources_matched: int = 0
    render_resources_unmatched: int = 0
    render_failed_requests: int = 0
    render_unloaded_images: int = 0
    images_skipped: int = 0
    accessible_images: int = 0
    errors: list[dict[str, str]] = field(default_factory=list)


@dataclass
class ResourceMatch:
    source: str
    kind: str
    local_path: str | None = None
    status: str = "unmatched"
    reason: str = ""


@dataclass
class RenderHtmlResult:
    raw_html: str
    prepared_html: str
    resources: list[ResourceMatch] = field(default_factory=list)
    console_messages: list[str] = field(default_factory=list)
    page_errors: list[str] = field(default_factory=list)
    failed_requests: list[str] = field(default_factory=list)
    unloaded_images: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def matched_count(self) -> int:
        return sum(1 for item in self.resources if item.status == "matched")

    @property
    def unmatched_count(self) -> int:
        return sum(1 for item in self.resources if item.status != "matched")


def parse_qid_filter(raw_qids: str | None, qid_file: str | None) -> set[str] | None:
    qids: list[str] = []
    if raw_qids:
        qids.extend(item.strip() for item in raw_qids.split(",") if item.strip())
    if qid_file:
        qids.extend(item.strip() for item in Path(qid_file).read_text(encoding="utf-8").splitlines() if item.strip())
    return set(qids) if qids else None


def load_tasks(path: Path, limit: int | None, qid_filter: set[str] | None = None) -> list[dict[str, Any]]:
    tasks: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as file:
        for line_number, line in enumerate(file, start=1):
            if not line.strip():
                continue
            task = json.loads(line)
            if qid_filter is not None and str(task.get("qid", "")) not in qid_filter:
                continue
            task["_line_number"] = line_number
            tasks.append(task)
            if limit and len(tasks) >= limit:
                break
    return tasks


def read_all_jsonl_count(path: Path) -> int:
    count = 0
    with path.open("r", encoding="utf-8") as file:
        for line in file:
            if line.strip():
                count += 1
    return count


def count_csv(path: Path) -> int | None:
    if not path.exists():
        return None
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return sum(1 for _ in csv.DictReader(file))


def count_sqlite(path: Path) -> int | None:
    if not path.exists():
        return None
    conn = sqlite3.connect(path)
    try:
        return int(conn.execute("SELECT COUNT(*) FROM tasks").fetchone()[0])
    finally:
        conn.close()


def parse_validation_status(report_path: Path) -> str:
    if not report_path.exists():
        return "not found"
    text = report_path.read_text(encoding="utf-8", errors="replace")
    return "OK" if "Status: OK" in text or "validation_status | OK" in text else "unknown"


def resolve_path(raw_path: str, data_path: Path) -> Path | None:
    candidate = Path(raw_path)
    candidates = [candidate]
    if not candidate.is_absolute():
        candidates.extend(
            [
                Path.cwd() / candidate,
                data_path.parent / candidate,
                data_path.parent.parent / candidate,
            ]
        )
    for item in candidates:
        if item.exists():
            return item.resolve()
    return None


def rendered_dir_from_args(args: argparse.Namespace, data_path: Path) -> Path:
    if args.render_dir:
        return Path(args.render_dir)
    return data_path.parent / "rendered_tasks"


def render_debug_dir_from_args(args: argparse.Namespace, data_path: Path) -> Path:
    return Path(args.render_debug_dir) if args.render_debug_dir else data_path.parent / "render_debug"


def image_uri(path: Path) -> str:
    mime_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def replacement_image_html(local_path: Path | None, fallback_src: str, alt: str = "") -> str:
    src = image_uri(local_path) if local_path is not None else fallback_src
    return f'<img src="{src}" alt="{alt}" class="inline-image">'


def normalize_resource_ref(value: str) -> str:
    value = unescape(unquote(str(value or "").strip().strip("\"'")))
    parsed = urlparse(value)
    path = parsed.path if parsed.scheme else value
    path = path.replace("\\", "/")
    while path.startswith("../"):
        path = path[3:]
    return path.lstrip("./").lower()


def resource_basename(value: str) -> str:
    return Path(urlparse(normalize_resource_ref(value)).path).name.lower()


def resource_stem(value: str) -> str:
    return Path(resource_basename(value)).stem.lower()


def looks_like_image_ref(value: str) -> bool:
    normalized = normalize_resource_ref(value)
    basename = resource_basename(normalized)
    return bool(
        re.search(r"\.(?:png|jpe?g|gif|webp|bmp|svg)(?:$|\?)", normalized, flags=re.IGNORECASE)
        or "docs/" in normalized
        or "questions/" in normalized
        or basename.startswith("innerimg")
    )


def string_args(js_args: str) -> list[str]:
    return [unescape(match.group(2)) for match in re.finditer(r"""(["'])(.*?)(?<!\\)\1""", js_args, flags=re.DOTALL)]


def extract_document_write_html(raw_html: str) -> str:
    def replace(match: re.Match[str]) -> str:
        args = string_args(match.group(1))
        fragments = [item for item in args if "<img" in item.lower()]
        return unescape("".join(fragments)) if fragments else match.group(0)

    return re.sub(r"document\.write\s*\((.*?)\)\s*;?", replace, raw_html, flags=re.IGNORECASE | re.DOTALL)


def extract_js_image_calls(raw_html: str) -> list[str]:
    refs: list[str] = []
    for match in re.finditer(r"\bshowpictureq?\s*\((.*?)\)", raw_html, flags=re.IGNORECASE | re.DOTALL):
        refs.extend(item for item in string_args(match.group(1)) if looks_like_image_ref(item))
    return refs


def extract_css_background_images(raw_html: str) -> list[str]:
    refs = []
    for match in re.finditer(r"background(?:-image)?\s*:\s*url\(([^)]+)\)", raw_html, flags=re.IGNORECASE):
        ref = match.group(1).strip().strip("\"'")
        if looks_like_image_ref(ref):
            refs.append(ref)
    return refs


def extract_img_sources(raw_html: str) -> list[str]:
    if lxml_html is None:
        return [
            unescape(src)
            for src in re.findall(r"<img\b[^>]*\bsrc\s*=\s*['\"]([^'\"]+)['\"]", raw_html, flags=re.IGNORECASE)
            if src
        ]
    try:
        fragment = lxml_html.fromstring(f"<div>{raw_html}</div>")
    except Exception:
        return [
            unescape(src)
            for src in re.findall(r"<img\b[^>]*\bsrc\s*=\s*['\"]([^'\"]+)['\"]", raw_html, flags=re.IGNORECASE)
            if src
        ]
    return [unescape(src) for src in fragment.xpath(".//img/@src") if src]


def ordered_unique(values: list[str]) -> list[str]:
    seen = set()
    result = []
    for value in values:
        key = normalize_resource_ref(value)
        if key not in seen:
            seen.add(key)
            result.append(value)
    return result


def extract_render_resources(raw_html: str) -> list[ResourceMatch]:
    expanded = extract_document_write_html(raw_html)
    refs = ordered_unique(
        extract_js_image_calls(expanded) + extract_img_sources(expanded) + extract_css_background_images(expanded)
    )
    return [ResourceMatch(source=ref, kind="image") for ref in refs]


def resource_keys(value: str, source_url: str | None = None) -> set[str]:
    values = {value}
    if source_url:
        values.add(urljoin(source_url, value))
    keys: set[str] = set()
    for item in values:
        normalized = normalize_resource_ref(item)
        basename = resource_basename(item)
        stem = resource_stem(item)
        for key in (normalized, basename, stem):
            if key:
                keys.add(key)
    return keys


def build_local_image_index(task: dict[str, Any], data_path: Path) -> tuple[dict[str, Path], list[Path]]:
    index: dict[str, Path] = {}
    ordered_paths: list[Path] = []
    source_url = str(task.get("source_url") or SOURCE_URL)
    refs = [str(item) for item in (task.get("image_refs") or [])]
    locals_ = [str(item) for item in (task.get("local_images") or [])]
    for raw_local in locals_:
        local_path = resolve_path(raw_local, data_path)
        if local_path is None:
            continue
        ordered_paths.append(local_path)
        for key in resource_keys(raw_local) | {local_path.name.lower(), local_path.stem.lower()}:
            if key:
                index[key] = local_path
    for ref, raw_local in zip(refs, locals_):
        local_path = resolve_path(raw_local, data_path)
        if local_path is None:
            continue
        for key in resource_keys(ref, source_url):
            if key:
                index[key] = local_path
    for idx, local_path in enumerate(ordered_paths):
        index[f"{str(task.get('qid', '')).lower()}:{idx}"] = local_path
    images_dir = data_path.parent / "images"
    if images_dir.exists():
        for local_path in images_dir.iterdir():
            if local_path.is_file():
                index.setdefault(local_path.name.lower(), local_path)
                index.setdefault(local_path.stem.lower(), local_path)
    return index, ordered_paths


def match_resources(task: dict[str, Any], data_path: Path, resources: list[ResourceMatch]) -> None:
    image_index, ordered_paths = build_local_image_index(task, data_path)
    for idx, resource in enumerate(resources):
        keys = resource_keys(resource.source, str(task.get("source_url") or SOURCE_URL))
        local_path = next((image_index[key] for key in keys if key in image_index), None)
        if local_path is None:
            local_path = image_index.get(f"{str(task.get('qid', '')).lower()}:{idx}")
        if local_path is None and idx < len(ordered_paths):
            local_path = ordered_paths[idx]
            resource.reason = "matched_by_order"
        if local_path is not None:
            resource.local_path = str(local_path)
            resource.status = "matched"
        else:
            resource.status = "unmatched"
            resource.reason = "no local image match"


def find_resource_for_ref(ref: str, resources: list[ResourceMatch]) -> ResourceMatch | None:
    ref_keys = resource_keys(ref, SOURCE_URL)
    for resource in resources:
        if ref_keys & resource_keys(resource.source, SOURCE_URL):
            return resource
    return None


def replace_js_image_calls(raw_html: str, resources: list[ResourceMatch]) -> str:

    def replacement(match: re.Match[str]) -> str:
        candidates = [item for item in string_args(match.group(1)) if looks_like_image_ref(item)]
        if not candidates:
            return match.group(0)
        resource = None
        for candidate in candidates:
            resource = find_resource_for_ref(candidate, resources)
            if resource is not None:
                break
        if resource is None:
            return match.group(0)
        local_path = Path(resource.local_path) if resource.local_path else None
        return replacement_image_html(local_path, urljoin(SOURCE_URL, resource.source), resource.source)

    html = re.sub(r"\bshowpictureq?\s*\((.*?)\)\s*;?", replacement, raw_html, flags=re.IGNORECASE | re.DOTALL)
    html = re.sub(
        r"<script\b[^>]*>\s*(<img\b[^>]*>)\s*</script>",
        lambda match: match.group(1),
        html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    html = re.sub(
        r"<script\b[^>]*>\s*document\.write\s*\(\s*(<img\b[^>]*>)\s*\)\s*;?\s*</script>",
        lambda match: match.group(1),
        html,
        flags=re.IGNORECASE | re.DOTALL,
    )
    return html


def replace_img_src_and_backgrounds(raw_html: str, resources: list[ResourceMatch]) -> str:
    matched_by_source = {normalize_resource_ref(item.source): item for item in resources}
    matched_by_basename = {resource_basename(item.source): item for item in resources if resource_basename(item.source)}

    def replacement_src(match: re.Match[str]) -> str:
        src = match.group(2)
        resource = matched_by_source.get(normalize_resource_ref(src)) or matched_by_basename.get(resource_basename(src))
        if resource and resource.local_path:
            return f'{match.group(1)}{image_uri(Path(resource.local_path))}{match.group(3)}'
        return match.group(0)

    html = re.sub(r"(\bsrc\s*=\s*['\"])([^'\"]+)(['\"])", replacement_src, raw_html, flags=re.IGNORECASE)

    def replacement_bg(match: re.Match[str]) -> str:
        src = match.group(1).strip().strip("\"'")
        resource = matched_by_source.get(normalize_resource_ref(src)) or matched_by_basename.get(resource_basename(src))
        if resource and resource.local_path:
            return f"url({image_uri(Path(resource.local_path))})"
        return match.group(0)

    return re.sub(r"url\(([^)]+)\)", replacement_bg, html, flags=re.IGNORECASE)


def prepare_render_html(task: dict[str, Any], data_path: Path) -> RenderHtmlResult:
    raw_html = str(task.get("html") or "")
    result = RenderHtmlResult(raw_html=raw_html, prepared_html="")
    if not raw_html.strip():
        result.errors.append("Source HTML is empty")
        return result

    expanded = extract_document_write_html(raw_html)
    resources = extract_render_resources(expanded)
    match_resources(task, data_path, resources)
    html = replace_js_image_calls(expanded, resources)
    html = replace_img_src_and_backgrounds(html, resources)
    result.resources = resources
    result.prepared_html = html
    return result


def create_rendered_html_document(task: dict[str, Any], data_path: Path, render_width: int) -> RenderHtmlResult:
    result = prepare_render_html(task, data_path)
    task_html = result.prepared_html
    if not task_html:
        return result
    result.prepared_html = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    padding: 24px;
    background: #ffffff;
    color: #111827;
    font-family: Arial, "Helvetica Neue", sans-serif;
    font-size: 18px;
    line-height: 1.45;
  }}
  .task-card {{
    width: {render_width}px;
    background: #ffffff;
    border: 1px solid #d1d5db;
    padding: 18px 20px;
  }}
  .hint {{
    margin: 0 0 12px;
    color: #374151;
    font-weight: 700;
  }}
  table {{
    border-collapse: collapse;
    max-width: 100%;
  }}
  td {{
    vertical-align: top;
  }}
  p {{
    margin: 0 0 10px;
  }}
  img,
  .inline-image {{
    max-width: 100%;
    height: auto;
    vertical-align: middle;
  }}
  input[type="text"] {{
    min-height: 26px;
    border: 1px solid #9ca3af;
  }}
  .submit-block {{
    display: none;
  }}
</style>
</head>
<body>
<div class="task-card">
{task_html}
</div>
</body>
</html>"""
    return result


class HtmlTaskRenderer:
    def __init__(self, render_width: int, render_scale: float) -> None:
        self.render_width = render_width
        self.render_scale = render_scale
        self.playwright = None
        self.browser = None
        self.available = False
        self.error = ""

    def __enter__(self) -> "HtmlTaskRenderer":
        try:
            from playwright.sync_api import sync_playwright

            self.playwright = sync_playwright().start()
            self.browser = self.playwright.chromium.launch(headless=True)
            self.available = True
        except Exception as exc:
            self.error = str(exc)
            self.close()
        return self

    def __exit__(self, exc_type, exc, traceback) -> None:
        self.close()

    def close(self) -> None:
        if self.browser is not None:
            self.browser.close()
            self.browser = None
        if self.playwright is not None:
            self.playwright.stop()
            self.playwright = None

    def render(self, html: str, out_path: Path, result: RenderHtmlResult | None = None) -> Path:
        if not self.available or self.browser is None:
            raise RuntimeError(self.error or "Playwright renderer is unavailable")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        page = self.browser.new_page(
            viewport={"width": self.render_width + 80, "height": 1400},
            device_scale_factor=self.render_scale,
        )
        try:
            if result is not None:
                page.on("console", lambda msg: result.console_messages.append(f"{msg.type}: {msg.text}"))
                page.on("pageerror", lambda exc: result.page_errors.append(str(exc)))
                page.on("requestfailed", lambda request: result.failed_requests.append(f"{request.url}: {request.failure}"))
            page.set_content(html, wait_until="load")
            page.wait_for_load_state("networkidle", timeout=10000)
            unloaded = page.evaluate(
                """async () => {
                    const images = Array.from(document.images);
                    await Promise.all(images.map(img => {
                        if (img.complete) return Promise.resolve();
                        return new Promise(resolve => {
                            img.addEventListener('load', resolve, { once: true });
                            img.addEventListener('error', resolve, { once: true });
                            setTimeout(resolve, 3000);
                        });
                    }));
                    return images
                        .filter(img => !img.complete || img.naturalWidth === 0)
                        .map(img => img.currentSrc || img.src || img.alt || '');
                }"""
            )
            if result is not None:
                result.unloaded_images.extend([str(item) for item in unloaded])
            page.locator(".task-card").screenshot(path=str(out_path))
        finally:
            page.close()
        return out_path


def safe_text(text: Any, stats: ExportStats, qid: str, field: str) -> str:
    value = "" if text is None else str(text)
    if len(value) > MAX_CELL_TEXT:
        stats.errors.append(
            {"qid": qid, "type": "too_long_text", "message": f"{field} truncated from {len(value)} chars"}
        )
        return value[:MAX_CELL_TEXT]
    return value


def options_text(options: list[dict[str, Any]]) -> str:
    lines = []
    for option in options or []:
        value = str(option.get("value", "")).strip()
        text = str(option.get("text", "")).strip()
        if value and text and value != text:
            lines.append(f"{value}: {text}")
        else:
            lines.append(text or value)
    return "\n".join(lines)


def normalized_text(value: Any) -> str:
    return " ".join(str(value or "").lower().replace("ё", "е").split())


def has_answer_options(task: dict[str, Any]) -> bool:
    for key in ("options", "choices", "answers"):
        value = task.get(key)
        if isinstance(value, list) and len(value) > 0:
            return True
    return False


def infer_task_type(task: dict[str, Any]) -> str:
    text = normalized_text(task.get("text", ""))
    if has_answer_options(task):
        return "Выбор варианта ответа"
    if any(
        phrase in text
        for phrase in (
            "впишите правильный ответ",
            "запишите ответ",
            "в ответе укажите",
            "ответ:",
            "найдите",
            "вычислите",
        )
    ):
        return "Краткий ответ"
    if any(
        phrase in text
        for phrase in (
            "укажите номера",
            "выберите все",
            "выберите верные",
            "какие из следующих",
            "несколько вариантов",
            "все верные утверждения",
            "верные утверждения",
        )
    ):
        return "Выбор нескольких ответов"
    if any(
        phrase in text
        for phrase in (
            "установите соответствие",
            "соотнесите",
            "каждому элементу",
            "подберите",
        )
    ):
        return "Установление соответствия"
    if any(phrase in text for phrase in ("расположите", "в правильном порядке", "последовательность")):
        return "Установление последовательности"
    return "Не определено"


def flatten_topic_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, list):
        parts = [flatten_topic_value(item) for item in value]
        return "; ".join(part for part in parts if part)
    if isinstance(value, dict):
        preferred = []
        for key in ("topic", "section", "theme", "subject", "block", "category", "kesh", "кэс", "breadcrumbs", "path"):
            item = flatten_topic_value(value.get(key))
            if item:
                preferred.append(item)
        if preferred:
            return "; ".join(dict.fromkeys(preferred))
    return ""


def explicit_topic(task: dict[str, Any]) -> str:
    for key in (
        "topic",
        "section",
        "theme",
        "subject",
        "block",
        "category",
        "kesh",
        "кэс",
        "kes",
        "metadata",
        "breadcrumbs",
        "path",
        "source section",
        "source_section",
    ):
        value = flatten_topic_value(task.get(key))
        if value:
            return value
    return ""


def contains_any(text: str, phrases: tuple[str, ...]) -> bool:
    return any(phrase in text for phrase in phrases)


def infer_math_topic(task: dict[str, Any]) -> str:
    value = explicit_topic(task)
    if value:
        return value

    text = normalized_text(task.get("text", ""))
    topic_rules = [
        ("Вероятность и статистика", ("вероятность", "случайный опыт", "элементарные события", "благоприятствуют событию", "среднее арифметическое", "медиана", "мода", "диаграмма", "таблица частот")),
        ("Функции и графики", ("график функции", "функция", "координатная плоскость", "парабола", "прямая", "гипербола", "значение функции", "область определения")),
        ("Уравнения и неравенства", ("система уравнений", "система неравенств", "решите уравнение", "корень уравнения", "неравенство")),
        ("Последовательности и прогрессии", ("арифметическая прогрессия", "геометрическая прогрессия", "последовательность", "n-й член", "сумма первых членов")),
        ("Геометрия: площади", ("площадь фигуры", "площадь треугольника", "площадь круга", "площадь трапеции", "площадь параллелограмма", "площадь")),
        ("Геометрия: окружность и круг", ("окружность", "круг", "радиус", "диаметр", "хорда", "дуга", "касательная", "центральный угол", "вписанный угол")),
        ("Геометрия: треугольники", ("подобные треугольники", "треугольник", "катет", "гипотенуза", "медиана", "биссектриса", "высота")),
        ("Геометрия: четырёхугольники и многоугольники", ("параллелограмм", "ромб", "трапеция", "прямоугольник", "квадрат", "многоугольник")),
        ("Геометрия: стереометрия", ("объем", "поверхность", "призма", "пирамида", "цилиндр", "конус", "шар", "куб", "параллелепипед")),
        ("Практико-ориентированные задачи", ("участок", "план", "квартира", "тариф", "квитанция", "печь", "дорога", "карта", "масштаб", "плитка", "теплица", "забор", "ремонт")),
        ("Текстовые задачи", ("поезд", "автомобиль", "скорость", "время", "расстояние", "работа", "производительность", "смесь", "сплав", "вклад", "цена", "стоимость", "покупка", "проценты по вкладу")),
        ("Комбинаторика", ("сколько способов", "варианты", "комбинации", "перестановки", "выбор", "код", "пароль")),
        ("Алгебраические выражения", ("упростите выражение", "преобразуйте выражение", "многочлен", "одночлен", "разложите на множители", "тождество")),
        ("Арифметика и вычисления", ("вычислите", "значение выражения", "дробь", "процент", "отношение", "пропорция", "округление", "степень", "корень")),
    ]
    for topic, phrases in topic_rules:
        if contains_any(text, phrases):
            return topic
    return "Не определено"


def set_common_layout(wb: Workbook) -> None:
    wb.properties.creator = "Codex"
    wb.properties.title = TITLE
    wb.properties.subject = "FIPI OGE math task export"


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in row:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def create_summary_sheet(
    wb: Workbook,
    data_path: Path,
    tasks: list[dict[str, Any]],
    total_jsonl_count: int,
    mode: str,
) -> None:
    ws = wb.active
    ws.title = SUMMARY_SHEET
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90

    values = [
        ("Название набора данных", TITLE),
        ("Источник", SOURCE_URL),
        ("Дата формирования Excel", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Количество заданий", len(tasks)),
        ("Количество изображений", sum(len(t.get("local_images") or []) for t in tasks)),
        ("Статус последней валидации", parse_validation_status(data_path.parent / "export_report.md")),
        ("Примечание о правах", "Права на задания и материалы сохраняются за ФИПИ/правообладателями."),
    ]
    ws["A1"] = "Параметр"
    ws["B1"] = "Значение"
    style_header(ws[1])
    for row_idx, (name, value) in enumerate(values, start=2):
        ws.cell(row_idx, 1, name)
        ws.cell(row_idx, 2, "" if value is None else value)
        ws.cell(row_idx, 1).font = Font(bold=True, color="374151")
        ws.cell(row_idx, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row_idx, 2).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def create_index_sheet(wb: Workbook, tasks: list[dict[str, Any]], data_path: Path, stats: ExportStats) -> None:
    ws = wb.create_sheet(INDEX_SHEET)
    headers = [
        "№",
        "Номер задания",
        "URL",
        "Тип задания",
        "Тема",
        "Текст задания",
        "Ссылка на карточку",
    ]
    ws.append(headers)
    style_header(ws[1])
    widths = [8, 18, 42, 28, 34, 90, 22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(tasks) + 1}"

    for row_idx, task in enumerate(tasks, start=2):
        task_no = row_idx - 1
        qid = str(task.get("qid", ""))
        task_type = infer_task_type(task)
        topic = infer_math_topic(task)
        full_text = safe_text(task.get("text", ""), stats, qid, "index_text")
        values = [
            task_no,
            task.get("task_number") or task_no,
            task.get("url", ""),
            task_type,
            topic,
            full_text,
            "Открыть карточку",
        ]
        ws.append(values)
        excel_row = row_idx
        task_anchor = task.get("_excel_task_row")
        if task_anchor:
            ws.cell(excel_row, 7).hyperlink = f"#{quote_sheetname(TASKS_SHEET)}!A{task_anchor}"
            ws.cell(excel_row, 7).style = "Hyperlink"
        ws.cell(excel_row, 3).hyperlink = task.get("url", "")
        ws.cell(excel_row, 3).style = "Hyperlink"
        for cell in ws[excel_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=len(tasks) + 1):
        text = str(ws.cell(row[0].row, 6).value or "")
        ws.row_dimensions[row[0].row].height = min(120, max(34, math.ceil(len(text) / 95) * 15))


def setup_tasks_sheet(wb: Workbook) -> Any:
    ws = wb.create_sheet(TASKS_SHEET)
    widths = [6, 14, 18, 18, 18, 18, 18, 18, 18, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A1"
    return ws


def merge_write(ws, row: int, start_col: int, end_col: int, value: str, style: str = "body") -> None:
    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    cell = ws.cell(row, start_col, value)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if style == "title":
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True, size=12, color="1F4E78")
    elif style == "service":
        cell.fill = PatternFill("solid", fgColor="F3F4F6")
        cell.font = Font(size=9, color="6B7280")
    else:
        cell.font = Font(size=11, color="111827")


def apply_card_border(ws, start_row: int, end_row: int, start_col: int = 1, end_col: int = 10) -> None:
    side = Side(style="thin", color="D1D5DB")
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=side if col == start_col else cell.border.left,
                right=side if col == end_col else cell.border.right,
                top=side if row == start_row else cell.border.top,
                bottom=side if row == end_row else cell.border.bottom,
            )


def image_size(path: Path) -> tuple[int, int]:
    with PILImage.open(path) as image:
        return image.size


def write_render_debug(
    debug_dir: Path,
    task: dict[str, Any],
    sequence: int,
    render_result: RenderHtmlResult,
    rendered_path: Path | None,
) -> None:
    qid = str(task.get("qid", ""))
    task_dir = debug_dir / f"{sequence:04d}_{qid}"
    task_dir.mkdir(parents=True, exist_ok=True)
    (task_dir / "source.html").write_text(render_result.raw_html, encoding="utf-8")
    (task_dir / "prepared.html").write_text(render_result.prepared_html, encoding="utf-8")
    debug_png = None
    if rendered_path is not None and rendered_path.exists():
        debug_png = task_dir / "rendered.png"
        shutil.copy2(rendered_path, debug_png)
    diagnostics = {
        "qid": qid,
        "url": task.get("url"),
        "line_number": task.get("_line_number"),
        "html_fields": [key for key, value in task.items() if "html" in str(key).lower() and value],
        "text_fields": [key for key, value in task.items() if "text" in str(key).lower() and value],
        "image_fields": [key for key, value in task.items() if "image" in str(key).lower() and value],
        "found_resources": [item.source for item in render_result.resources],
        "matched_resources": [
            {"source": item.source, "local_path": item.local_path, "reason": item.reason}
            for item in render_result.resources
            if item.status == "matched"
        ],
        "unmatched_resources": [
            {"source": item.source, "reason": item.reason}
            for item in render_result.resources
            if item.status != "matched"
        ],
        "console_messages": render_result.console_messages,
        "page_errors": render_result.page_errors,
        "failed_requests": render_result.failed_requests,
        "unloaded_images": render_result.unloaded_images,
        "errors": render_result.errors,
        "rendered_png": str(debug_png or rendered_path) if rendered_path else None,
    }
    (task_dir / "diagnostics.json").write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")


def scaled_size(width: int, height: int, max_width: int = MAX_IMAGE_WIDTH_PX) -> tuple[int, int]:
    if width <= max_width:
        return width, height
    scale = max_width / width
    return int(width * scale), int(height * scale)


def build_tasks_sheet(
    ws,
    tasks: list[dict[str, Any]],
    data_path: Path,
    mode: str,
    stats: ExportStats,
    task_view: str,
    renderer: HtmlTaskRenderer | None,
    render_dir: Path,
    render_width: int,
    debug_render: bool,
    render_debug_dir: Path,
) -> None:
    iterable = tasks
    if tqdm:
        iterable = tqdm(tasks, desc=f"Excel {mode}", unit="task")
    row = 1
    for idx, task in enumerate(iterable, start=1):
        qid = str(task.get("qid", ""))
        card_start = row
        task["_excel_task_row"] = row
        try:
            text = safe_text(task.get("text", ""), stats, qid, "text")
            if not text.strip():
                stats.errors.append({"qid": qid, "type": "empty_text", "message": "Task text is empty"})
            title = f"Задание № {idx} / {qid} / источник ФИПИ"
            merge_write(ws, row, 1, 10, title, "title")
            ws.cell(row, 1).hyperlink = task.get("url", "")
            ws.cell(row, 1).font = Font(bold=True, size=12, color="0563C1", underline="single")
            ws.cell(row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
            ws.row_dimensions[row].height = 26
            row += 1

            rendered_done = False
            if task_view == "rendered":
                render_result = create_rendered_html_document(task, data_path, render_width)
                html_doc = render_result.prepared_html
                rendered_path = render_dir / f"{idx:04d}_{qid}.png"
                stats.render_resources_found += len(render_result.resources)
                stats.render_resources_matched += render_result.matched_count
                stats.render_resources_unmatched += render_result.unmatched_count
                for item in render_result.resources:
                    if item.status != "matched":
                        stats.errors.append({"qid": qid, "type": "unmatched_render_resource", "message": item.source})
                if not html_doc:
                    stats.rendered_fallbacks += 1
                    stats.errors.append(
                        {"qid": qid, "type": "rendered_fallback", "message": "; ".join(render_result.errors) or "Source HTML is empty"}
                    )
                elif renderer is None or not renderer.available:
                    stats.rendered_fallbacks += 1
                    message = renderer.error if renderer is not None and renderer.error else "Playwright renderer is unavailable"
                    stats.errors.append({"qid": qid, "type": "rendered_fallback", "message": message})
                else:
                    try:
                        renderer.render(html_doc, rendered_path, render_result)
                        stats.render_failed_requests += len(render_result.failed_requests)
                        stats.render_unloaded_images += len(render_result.unloaded_images)
                        for failed in render_result.failed_requests:
                            stats.errors.append({"qid": qid, "type": "render_failed_request", "message": failed})
                        for unloaded in render_result.unloaded_images:
                            stats.errors.append({"qid": qid, "type": "render_unloaded_image", "message": unloaded})
                        stats.rendered_png_created += 1
                        width, height = image_size(rendered_path)
                        scaled_w, scaled_h = scaled_size(width, height)
                        if mode == "embedded":
                            img = XLImage(str(rendered_path))
                            img.width = scaled_w
                            img.height = scaled_h
                            ws.add_image(img, f"B{row}")
                            ws.row_dimensions[row].height = max(28, scaled_h * 0.75 + 10)
                            stats.images_embedded += 1
                        else:
                            merge_write(ws, row, 1, 10, "Открыть rendered-карточку", "service")
                            ws.cell(row, 1).hyperlink = rendered_path.resolve().as_uri()
                            ws.cell(row, 1).style = "Hyperlink"
                            ws.row_dimensions[row].height = 24
                            stats.images_linked += 1
                        row += 1
                        rendered_done = True
                    except Exception as exc:
                        stats.rendered_fallbacks += 1
                        stats.errors.append({"qid": qid, "type": "rendered_fallback", "message": str(exc)})
                if debug_render:
                    write_render_debug(render_debug_dir, task, idx, render_result, rendered_path if rendered_path.exists() else None)

            if rendered_done:
                apply_card_border(ws, card_start, row - 1)
                row += 2
                stats.cards_created += 1
                stats.tasks_processed += 1
                continue

            merge_write(ws, row, 1, 10, text)
            ws.row_dimensions[row].height = min(260, max(45, math.ceil(len(text) / 95) * 16))
            row += 1

            opt_text = options_text(task.get("options") or [])
            if opt_text:
                opt_text = safe_text("Варианты ответа:\n" + opt_text, stats, qid, "options")
                merge_write(ws, row, 1, 10, opt_text)
                ws.row_dimensions[row].height = min(220, max(35, math.ceil(len(opt_text) / 90) * 16))
                row += 1

            local_images = task.get("local_images") or []
            if local_images:
                merge_write(ws, row, 1, 10, "Изображения / схемы / графики", "service")
                ws.row_dimensions[row].height = 20
                row += 1
            for img_idx, raw_path in enumerate(local_images, start=1):
                image_path = resolve_path(str(raw_path), data_path)
                if image_path is None:
                    stats.images_skipped += 1
                    stats.errors.append({"qid": qid, "type": "missing_image", "message": str(raw_path)})
                    merge_write(ws, row, 1, 10, f"Изображение не найдено: {raw_path}", "service")
                    row += 1
                    continue
                stats.accessible_images += 1
                if mode == "links":
                    merge_write(ws, row, 1, 10, f"Открыть изображение {img_idx}", "service")
                    ws.cell(row, 1).hyperlink = image_path.as_uri()
                    ws.cell(row, 1).style = "Hyperlink"
                    ws.row_dimensions[row].height = 22
                    stats.images_linked += 1
                    row += 1
                    continue
                try:
                    width, height = image_size(image_path)
                    scaled_w, scaled_h = scaled_size(width, height)
                    img = XLImage(str(image_path))
                    img.width = scaled_w
                    img.height = scaled_h
                    ws.add_image(img, f"B{row}")
                    ws.row_dimensions[row].height = max(24, scaled_h * 0.75 + 8)
                    stats.images_embedded += 1
                    row += 1
                except Exception as exc:
                    stats.images_skipped += 1
                    stats.errors.append({"qid": qid, "type": "image_read_error", "message": f"{image_path}: {exc}"})
                    merge_write(ws, row, 1, 10, f"Не удалось встроить изображение: {image_path}", "service")
                    row += 1

            apply_card_border(ws, card_start, row - 1)
            row += 2
            stats.cards_created += 1
            stats.tasks_processed += 1
        except Exception as exc:
            stats.errors.append({"qid": qid, "type": "task_error", "message": str(exc)})
            row = max(row + 2, card_start + 3)


def create_errors_sheet(wb: Workbook, errors: list[dict[str, str]]) -> None:
    if not errors:
        return
    ws = wb.create_sheet(ERRORS_SHEET)
    ws.append(["qid", "type", "message"])
    style_header(ws[1])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 100
    for error in errors:
        ws.append([error.get("qid"), error.get("type"), error.get("message")])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(errors) + 1}"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def order_sheets(wb: Workbook) -> None:
    desired = [SUMMARY_SHEET, INDEX_SHEET, TASKS_SHEET, ERRORS_SHEET]
    wb._sheets.sort(key=lambda ws: desired.index(ws.title) if ws.title in desired else len(desired))


def validate_workbook(
    path: Path,
    tasks: list[dict[str, Any]],
    stats: ExportStats,
    mode: str,
    task_view: str,
) -> list[str]:
    issues: list[str] = []
    wb = load_workbook(path, read_only=False)
    try:
        expected_sheets = [SUMMARY_SHEET, INDEX_SHEET, TASKS_SHEET]
        if wb.sheetnames[:3] != expected_sheets:
            issues.append(f"Sheet order mismatch: {wb.sheetnames[:3]} != {expected_sheets}")
        index = wb[INDEX_SHEET]
        tasks_sheet = wb[TASKS_SHEET]
        index_rows = max(0, index.max_row - 1)
        if index_rows != len(tasks):
            issues.append(f"Index row count mismatch: {index_rows} != {len(tasks)}")
        qids = [str(task.get("qid", "")) for task in tasks]
        if len(set(qids)) != len(qids):
            issues.append("Duplicate qids found in source tasks")
        card_count = 0
        for row in range(1, tasks_sheet.max_row + 1):
            cell = tasks_sheet.cell(row=row, column=1)
            value = str(cell.value or "")
            fill = cell.fill.fgColor.rgb
            if value.startswith("Задание") and fill in {"00D9EAF7", "FFD9EAF7"}:
                card_count += 1
        if card_count != len(tasks):
            issues.append(f"Task card count mismatch: {card_count} != {len(tasks)}")
        for task in tasks:
            for raw_path in task.get("local_images") or []:
                if resolve_path(str(raw_path), Path(path).parent / "tasks.jsonl") is None:
                    issues.append(f"Missing local image link: {raw_path}")
        if mode == "embedded" and task_view == "structured" and stats.images_embedded != stats.accessible_images:
            issues.append(f"Embedded image mismatch: {stats.images_embedded} != {stats.accessible_images}")
        if mode == "embedded" and task_view == "rendered" and stats.rendered_png_created + stats.rendered_fallbacks != len(tasks):
            issues.append(
                "Rendered task count mismatch: "
                f"{stats.rendered_png_created} rendered + {stats.rendered_fallbacks} fallback != {len(tasks)}"
            )
    finally:
        wb.close()
    return issues


def build_excel(args: argparse.Namespace) -> ExportStats:
    data_path = Path(args.data)
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    qid_filter = parse_qid_filter(args.qid, args.qid_file)
    tasks = load_tasks(data_path, args.limit, qid_filter)
    total_jsonl_count = read_all_jsonl_count(data_path)
    render_dir = rendered_dir_from_args(args, data_path)
    render_debug_dir = render_debug_dir_from_args(args, data_path)
    stats = ExportStats()

    wb = Workbook()
    set_common_layout(wb)
    create_summary_sheet(wb, data_path, tasks, total_jsonl_count, args.mode)
    tasks_sheet = setup_tasks_sheet(wb)

    if args.task_view == "rendered":
        with HtmlTaskRenderer(args.render_width, args.render_scale) as renderer:
            build_tasks_sheet(
                tasks_sheet,
                tasks,
                data_path,
                args.mode,
                stats,
                args.task_view,
                renderer,
                render_dir,
                args.render_width,
                args.debug_render,
                render_debug_dir,
            )
    else:
        build_tasks_sheet(
            tasks_sheet,
            tasks,
            data_path,
            args.mode,
            stats,
            args.task_view,
            None,
            render_dir,
            args.render_width,
            args.debug_render,
            render_debug_dir,
        )
    create_index_sheet(wb, tasks, data_path, stats)
    create_errors_sheet(wb, stats.errors)
    order_sheets(wb)
    wb.save(out_path)

    issues = validate_workbook(out_path, tasks, stats, args.mode, args.task_view)
    for issue in issues:
        stats.errors.append({"qid": "", "type": "workbook_validation", "message": issue})
    if issues:
        # Persist validation issues to the workbook after initial save.
        wb = load_workbook(out_path)
        if ERRORS_SHEET in wb.sheetnames:
            ws = wb[ERRORS_SHEET]
        else:
            ws = wb.create_sheet(ERRORS_SHEET)
            ws.append(["qid", "type", "message"])
            style_header(ws[1])
        for issue in issues:
            ws.append(["", "workbook_validation", issue])
        order_sheets(wb)
        wb.save(out_path)
        wb.close()
    return stats


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create formatted Excel workbooks from FIPI JSONL export.")
    parser.add_argument("--data", required=True, help="Path to tasks.jsonl.")
    parser.add_argument("--out", required=True, help="Output .xlsx path.")
    parser.add_argument("--mode", choices=("embedded", "links"), default="links", help="Embed images or link to them.")
    parser.add_argument(
        "--task-view",
        choices=("structured", "rendered"),
        default="structured",
        help="Show tasks as editable text blocks or rendered HTML screenshots.",
    )
    parser.add_argument("--render-width", type=int, default=900, help="Rendered task card width in pixels.")
    parser.add_argument("--render-scale", type=float, default=1.0, help="Playwright device scale factor for screenshots.")
    parser.add_argument("--render-dir", default=None, help="Directory for rendered task PNG files.")
    parser.add_argument("--debug-render", action="store_true", help="Save rendered HTML/PNG diagnostics per task.")
    parser.add_argument("--render-debug-dir", default=None, help="Directory for rendered diagnostics.")
    parser.add_argument("--qid", default=None, help="Comma-separated qid filter.")
    parser.add_argument("--qid-file", default=None, help="Text file with one qid per line.")
    parser.add_argument("--limit", type=int, default=None, help="Limit number of tasks for preview.")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    stats = build_excel(args)
    out_path = Path(args.out)
    size_mb = out_path.stat().st_size / (1024 * 1024) if out_path.exists() else 0
    print("Excel export report")
    print(f"  tasks processed: {stats.tasks_processed}")
    print(f"  cards created: {stats.cards_created}")
    print(f"  images embedded: {stats.images_embedded}")
    print(f"  images linked: {stats.images_linked}")
    print(f"  rendered png created: {stats.rendered_png_created}")
    print(f"  rendered fallbacks: {stats.rendered_fallbacks}")
    print(f"  render resources found: {stats.render_resources_found}")
    print(f"  render resources matched: {stats.render_resources_matched}")
    print(f"  render resources unmatched: {stats.render_resources_unmatched}")
    print(f"  render failed requests: {stats.render_failed_requests}")
    print(f"  render unloaded images: {stats.render_unloaded_images}")
    print(f"  images skipped: {stats.images_skipped}")
    print(f"  errors: {len(stats.errors)}")
    print(f"  xlsx size: {size_mb:.2f} MB")
    if args.mode == "embedded" and size_mb > 200:
        print("  warning: embedded workbook is large and may open slowly")
    return 0 if not stats.errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
